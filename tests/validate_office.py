from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook


OUTPUT = Path(__file__).parent / "output"
PAYLOAD = "=2+2 <bolts> & washers"


def validate_zip(path: Path, required_names: set[str]) -> None:
    with ZipFile(path) as package:
        bad_member = package.testzip()
        assert bad_member is None, f"CRC failure in {bad_member}"
        assert required_names.issubset(set(package.namelist()))


def validate_xlsx() -> None:
    path = OUTPUT / "inventory.xlsx"
    validate_zip(path, {"[Content_Types].xml", "xl/workbook.xml", "xl/worksheets/sheet1.xml"})

    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Inventory"]
    assert sheet["A1"].value == "Item"
    assert sheet["B1"].value == "Quantity"
    assert sheet["A2"].value == PAYLOAD
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 7
    assert sheet["A3"].value == "Cable bundle"
    assert sheet["B3"].value == 12


def validate_docx() -> None:
    path = OUTPUT / "inventory.docx"
    validate_zip(path, {"[Content_Types].xml", "word/document.xml"})

    document = Document(path)
    assert document.paragraphs[0].text == "Inventory Report"
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "Item"
    assert table.cell(0, 1).text == "Quantity"
    assert table.cell(1, 0).text == PAYLOAD
    assert table.cell(1, 1).text == "7"
    assert table.cell(2, 0).text == "Cable bundle"
    assert table.cell(2, 1).text == "12"


if __name__ == "__main__":
    validate_xlsx()
    validate_docx()
    print("Office exports parsed successfully.")
